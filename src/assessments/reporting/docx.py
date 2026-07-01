"""Word document rendering for assessment health check deliverables."""

from __future__ import annotations

import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

from .context import HealthCheckReportContext


TEMPLATE_PATH = Path(__file__).resolve().parent / "docx" / "PAN Health Check Deliverable.docx"
SUMMARY_TABLE_INDEX = 3
FIREWALL_SECURITY_RULE_TABLE_INDEX = 16
RATING_STYLE_BY_LABEL = {
    "Critical": "Rating - Critical",
    "High": "Rating - High",
    "Medium": "Rating - Medium",
    "Low": "Rating - Low",
    "Informational": "Rating - Informational",
}
RATING_STYLE_BY_RISK = {
    "D-1": "Rating - Critical",
    "D-2": "Rating - High",
    "D-3": "Rating - Medium",
    "D-4": "Rating - Low",
    "D-5": "Rating - Informational",
}


def render_health_check_report(
    *,
    output_path: str | Path,
    context: HealthCheckReportContext,
    client_name: str,
    short_name: str,
    report_date: str,
    revision_number: str,
    opportunity_number: str,
    template_path: str | Path = TEMPLATE_PATH,
) -> Path:
    template = DocxTemplate(str(template_path))
    template.render(
        {
            "client_name": client_name,
            "short_name": short_name,
            "report_date": report_date,
            "revision_number": revision_number,
            "opportunity_number": opportunity_number,
        }
    )

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temp_file:
        temp_path = Path(temp_file.name)
    try:
        template.save(temp_path)
        document = Document(temp_path)
        _remove_environment_and_panorama_sections(document)
        _populate_summary_table(document, context)
        _replace_firewall_placeholder_tables(document, context)
        _replace_literal_text(document, "Firewall Y", "Reviewed Firewall Policy Scopes")
        document.save(output_path)
    finally:
        temp_path.unlink(missing_ok=True)

    return output_path


def _populate_summary_table(document: Document, context: HealthCheckReportContext) -> None:
    table = document.tables[SUMMARY_TABLE_INDEX]
    prototype_row = table.rows[1]
    while len(table.rows) > 1:
        _remove_row(table, 1)

    for row in context.summary_rows:
        new_row = _append_cloned_row(table, prototype_row)
        _set_cell_text(new_row.cells[0], "")
        _set_cell_text(new_row.cells[1], row.observation)
        _set_cell_text(new_row.cells[2], row.recommendation)
        _set_cell_text(
            new_row.cells[3],
            row.importance,
            paragraph_style=RATING_STYLE_BY_LABEL.get(row.importance),
        )


def _replace_firewall_placeholder_tables(document: Document, context: HealthCheckReportContext) -> None:
    blocks = list(_iter_body_blocks(document))
    removal_start = None
    removal_end = None
    prototype_caption = None
    prototype_table = None

    for idx, (kind, block) in enumerate(blocks):
        if kind == "p":
            text = block.text.strip()
            if text == "The following Security rules show risks identified for the firewall(s).":
                removal_start = idx
                if idx + 1 < len(blocks) and blocks[idx + 1][0] == "p":
                    prototype_caption = blocks[idx + 1][1]
                if idx + 2 < len(blocks) and blocks[idx + 2][0] == "t":
                    prototype_table = blocks[idx + 2][1]
            elif text == "Recommendations":
                removal_end = idx
                break

    if removal_start is None or removal_end is None or prototype_caption is None or prototype_table is None:
        prototype_table = document.tables[FIREWALL_SECURITY_RULE_TABLE_INDEX]
        _populate_firewall_security_rule_table(document, context, prototype_table)
        return

    for idx in range(removal_end - 1, removal_start - 1, -1):
        element = blocks[idx][1]._element
        element.getparent().remove(element)

    anchor = next(
        block._element
        for kind, block in _iter_body_blocks(document)
        if kind == "p" and block.text.strip() == "Recommendations"
    )
    table_number = 13
    for summary_row in context.summary_rows:
        _insert_cloned_paragraph_before(
            anchor_element=anchor,
            prototype_paragraph=prototype_caption,
            text=f"Table {table_number}: {summary_row.detail_table.title}",
        )
        _insert_findings_table_before(
            anchor_element=anchor,
            prototype_table=prototype_table,
            detail_table=summary_row.detail_table,
        )
        table_number += 1


def _populate_firewall_security_rule_table(document: Document, context: HealthCheckReportContext, table) -> None:
    prototype_row = table.rows[1]
    while len(table.rows) > 1:
        _remove_row(table, 1)

    for row in context.firewall_detail_rows:
        new_row = _append_cloned_row(table, prototype_row)
        _set_cell_text(new_row.cells[0], row.line_reference)
        _set_cell_text(new_row.cells[1], row.description)
        _set_cell_text(
            new_row.cells[2],
            row.risk_rating,
            paragraph_style=RATING_STYLE_BY_RISK.get(row.risk_rating),
        )


def _replace_literal_text(document: Document, old_text: str, new_text: str) -> None:
    for paragraph in document.paragraphs:
        if old_text in paragraph.text:
            paragraph.text = paragraph.text.replace(old_text, new_text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if old_text in paragraph.text:
                        paragraph.text = paragraph.text.replace(old_text, new_text)


def _remove_row(table, row_index: int) -> None:
    row = table.rows[row_index]
    row._element.getparent().remove(row._element)


def _remove_environment_and_panorama_sections(document: Document) -> None:
    blocks = list(_iter_body_blocks(document))
    start = None
    end = None

    for idx, (kind, block) in enumerate(blocks):
        if kind != "p":
            continue
        text = block.text.strip()
        if text == "Environment Review and Analysis":
            start = idx
        elif text == "PAN Firewall(s)":
            end = idx
            break

    if start is None or end is None or start >= end:
        return

    for idx in range(end - 1, start - 1, -1):
        element = blocks[idx][1]._element
        element.getparent().remove(element)


def _append_cloned_row(table, prototype_row):
    new_tr = deepcopy(prototype_row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def _set_cell_text(cell, text: str, *, paragraph_style: str | None = None) -> None:
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)
    if paragraph_style is not None:
        paragraph.style = paragraph_style
    paragraph.add_run(text)


def _iter_body_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield ("p", Paragraph(child, document))
        elif isinstance(child, CT_Tbl):
            yield ("t", Table(child, document))


def _insert_cloned_paragraph_before(anchor_element, *, prototype_paragraph: Paragraph, text: str) -> None:
    paragraph_element = deepcopy(prototype_paragraph._element)
    anchor_element.addprevious(paragraph_element)
    paragraph = Paragraph(paragraph_element, prototype_paragraph._parent)
    _set_paragraph_text(paragraph, text)


def _insert_findings_table_before(anchor_element, prototype_table, *, detail_table) -> None:
    table_element = deepcopy(prototype_table._tbl)
    anchor_element.addprevious(table_element)
    table = Table(table_element, prototype_table._parent)
    _set_cell_text(table.rows[0].cells[0], detail_table.title)
    prototype_row = prototype_table.rows[1]
    _set_cell_text(table.rows[1].cells[0], detail_table.first_column_heading)
    _set_cell_text(table.rows[1].cells[1], "Description")
    _set_cell_text(table.rows[1].cells[2], "Criticality/Severity")
    _set_row_as_heading(table.rows[1])
    while len(table.rows) > 2:
        _remove_row(table, 2)

    for detail_row in detail_table.detail_rows:
        new_row = _append_cloned_row(table, prototype_row)
        _set_cell_text(new_row.cells[0], f"{detail_row.primary}\n{detail_row.secondary}")
        _set_cell_text(new_row.cells[1], detail_row.description)
        _set_cell_text(
            new_row.cells[2],
            detail_row.severity_label,
            paragraph_style=RATING_STYLE_BY_LABEL.get(detail_row.severity_label),
        )


def _set_row_as_heading(row) -> None:
    for cell in row.cells:
        paragraph = cell.paragraphs[0]
        paragraph.style = "Table - Heading"
        for run in paragraph.runs:
            if hasattr(run.font, "bold"):
                run.font.bold = True
    if len(row.cells) >= 3:
        row.cells[2].paragraphs[0].alignment = 1


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)
    paragraph.add_run(text)
